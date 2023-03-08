# This program was created to assist in transforming raw survey data into a formatted csv that can be imported into a database

import openpyxl as xl
import csv
from pprint import pprint
from Insightly import insightly
import requests as r
from time import sleep


# Functions
def contact_finder(email_list):
    object_ids = []
    for email in email_list:
        try:
            search = r.get(
                f'https://api.insightly.com/v3.1/Contacts/Search?field_name=EMAIL_ADDRESS&field_value={email}',
                auth=(insightly.username, insightly.password))
            contact = search.json()[0]
            if contact:
                contact_id = contact['CONTACT_ID']
                object_ids.append(contact_id)
        except IndexError:
            try:
                search = r.get(f'https://api.insightly.com/v3.1/Leads/Search?field_name=EMAIL&field_value={email}',
                               auth=(insightly.username, insightly.password))
                lead = search.json()[0]
                if lead:
                    lead_id = lead['LEAD_ID']
                    object_ids.append(lead_id)
                else:
                    print(f'{email} could not be found')
            except IndexError:
                print(f'{email} could not be found')

    return object_ids


def column_finder(search_term):
    col = 1
    for cell in sheet[1]:
        if search_term not in str(cell.value).lower():
            col += 1
        else:
            break
    return col


def column_finder2(search_term):
    col = 1
    for cell in sheet[2]:
        if search_term not in str(cell.value).lower():
            col += 1
        else:
            break
    return col


# Program Start
print('Hi Alex! I\'m here to help with your accelerator surveys!! (っ´ω`c)♡')
print('Where is your excel file located? (P.S. Make sure it has a "Work Email:" header! and no merged cells')
file_path = input()

wb = xl.load_workbook(file_path)
sheet_names = wb.sheetnames
print('Which sheet are you working with?')
for index, value in enumerate(wb.sheetnames):
    print(f'{index} - {value}')
sheet_index = int(input())
sheet = wb[sheet_names[sheet_index]]
row_count = 1
for cell in sheet['A']:
    if cell.value is not None:
        row_count += 1

print('Is this for (P)rogress or (F)oundations?')
program = str(input()).lower()
cohort_options = insightly.get_custom_field_options('Accelerator_Survey__c', 'Cohort_Date__c')
if program == 'f':
    # Map columns
    csv_filename = 'Accelerator Foundations Pre-Survey Data Cleaned'
    cohort_col = column_finder('cohort')
    cohort_date_excel = str(sheet.cell(row=3, column=cohort_col).value)
    try:
        cohort_day = cohort_date_excel.split()[2]
        if cohort_day[0] == '0':
            cohort_day = cohort_day[1:]
        cohort_date = [x['OPTION_VALUE'] for x in cohort_options if cohort_day in x['OPTION_VALUE']][0]
        print('Just to double check, is this the correct cohort date? (y/n)')
        while True:
            print(cohort_date)
            forward = input()
            if str(forward).lower() == 'y':
                break
            elif str(forward).lower() == 'n':
                pass
    except IndexError:
        cohort_date = None
elif program == 'p':
    csv_filename = 'Accelerator Progress Pre-Survey Data Cleaned'
    for index, value in enumerate(cohort_options):
        print(index, value['OPTION_VALUE'])
    print('Select the cohort date for this Progress cohort')
    progress_index = int(input())
    cohort_date = cohort_options[progress_index]['OPTION_VALUE']


print('Thanks! Let me map the columns! One second...')
email_col = column_finder('work email')
age_col = column_finder('age')
gender_col = column_finder('gender')
race_col = column_finder('race')
race = []
caregiver_col = column_finder('caregiver')
caregiver = []
partner_col = column_finder('partner')
current_position_col = column_finder('current position')
current_company_col = column_finder('been at your current company')
promoted_col = column_finder('promoted')
engagement_col = column_finder2('engaged')
community_col = column_finder2('community')
retention_col = column_finder2('anticipate')

print('All mapped! Now I\'m combining values~')

# Combine Race columns into a single value
for row in range(3, row_count + 1):
    race_options = []
    for col in range(race_col, race_col + 8):
        cell_value = sheet.cell(row=row, column=col).value
        if cell_value is not None:
            race_options.append(cell_value)
    race_options_combined = ";".join(race_options)
    if race_options_combined == '':
        race_options_combined = None
    race.append(race_options_combined)

# Combine caregiver columns into a single value
for row in range(3, row_count + 1):
    caregiver_options = []
    for col in range(caregiver_col, caregiver_col + 6):
        cell_value = sheet.cell(row=row, column=col).value
        if cell_value is not None:
            caregiver_options.append(cell_value)
    caregiver_options_combined = ";".join(caregiver_options)
    if caregiver_options_combined == '':
        caregiver_options_combined = None
    caregiver.append(caregiver_options_combined)

print('Searching for Contact IDs~ Give me one second...')

while True:
    emails = []

    for cell in sheet[1]:
        if cell.value == 'Work Email:':
            email_column = str(cell)[-3]

    for cell in sheet[email_column]:
        if '@' in str(cell.value):
            emails.append(cell.value)
        else:
            pass

    contact_ids = contact_finder(emails)
    print(f'I found {len(contact_ids)} Contact IDs:')
    print(contact_ids)
    print('Should I search (a)gain or (c)ontinue?')
    proceed = str(input()).lower()
    if proceed == 'a':
        sleep(1)
        print('Okay, I\'ll try again...')
        pass
    elif proceed == 'c':
        break

print('Writing each row...')
rows = []
list_index = 0
row_index = 3

for contact_id in contact_ids:
    contact = insightly.get_object('Contact', contact_id)
    org = insightly.get_object('Organisation', contact['ORGANISATION_ID'])
    contact_full_name = contact['FIRST_NAME'] + ' ' + contact['LAST_NAME']

    row_dict = {'Accelerator Survey Name': '2023 - Foundations - ' + contact_full_name,
                 'Accelerator': str(contact_id) + ';' + contact_full_name,
                 'Title': contact['TITLE'],
                 'Organization': str(contact['ORGANISATION_ID']) + ';' + org['ORGANISATION_NAME'],
                 'Cohort Date': cohort_date,
                 'Age': sheet.cell(row=row_index, column=age_col).value,
                 'Gender Identity': sheet.cell(row=row_index, column=gender_col).value,
                 'Race/Ethnicity': race[list_index],
                 'Caregiver': caregiver[list_index],
                 'Partner Status': sheet.cell(row=row_index, column=partner_col).value,
                 'Years In Current Position': sheet.cell(row=row_index, column=current_position_col).value,
                 'Years At Current Company': sheet.cell(row=row_index, column=current_company_col).value,
                 'Last Promoted': sheet.cell(row=row_index, column=promoted_col).value,
                 'Engagement (PRE)': sheet.cell(row=row_index, column=engagement_col).value,
                 'Community (PRE)': sheet.cell(row=row_index, column=community_col).value,
                 'Retention (PRE)': sheet.cell(row=row_index, column=retention_col).value}

    rows.append(row_dict)
    list_index += 1
    row_index += 1

with open(f'/Users/alex/Downloads/{csv_filename}.csv', 'w', encoding='UTF8', newline='') as f:
    field_names = ['Accelerator Survey Name', 'Accelerator', 'Title', 'Organization', 'Cohort Date', 'Age',
                   'Gender Identity', 'Race/Ethnicity', 'Caregiver', 'Partner Status', 'Years In Current Position',
                   'Years At Current Company', 'Last Promoted', 'Engagement (PRE)', 'Community (PRE)',
                   'Retention (PRE)']
    writer = csv.DictWriter(f, fieldnames=field_names)
    writer.writeheader()
    writer.writerows(rows)

print(f'Done! You can find {csv_filename} in your Downloads folder')
